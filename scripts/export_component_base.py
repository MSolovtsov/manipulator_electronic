#!/usr/bin/env python3
"""Экспорт листов bom/Component_Base.xlsx в bom/csv/*.csv для просмотра диффов в Git.

Запускать после каждой правки базы:
    python3 scripts/export_component_base.py

Экспортируются вычисленные значения (не формулы), поэтому перед запуском файл
должен быть сохранён в Excel/LibreOffice (или пересчитан), иначе в ячейках с
формулами будут пустые значения.
"""
import csv
import sys
from pathlib import Path

try:
    import openpyxl
except ImportError:
    sys.exit("Нужен openpyxl: pip install openpyxl")

ROOT = Path(__file__).resolve().parents[1]
XLSX = ROOT / "bom" / "Component_Base.xlsx"
OUT = ROOT / "bom" / "csv"

# лист в базе -> имя CSV (латиницей — правило именования файлов репозитория)
SHEETS = {
    "Компоненты": "Components.csv",
    "Модули": "Modules.csv",
    "Отклонено": "Rejected.csv",
    "Сводка": "Summary.csv",
    "Справочники": "Lists.csv",
}


def cell(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v)


def main() -> int:
    if not XLSX.exists():
        sys.exit(f"Не найден {XLSX}")
    wb = openpyxl.load_workbook(XLSX, data_only=True)
    OUT.mkdir(parents=True, exist_ok=True)
    missing_values = 0
    for sheet, name in SHEETS.items():
        if sheet not in wb.sheetnames:
            print(f"! лист «{sheet}» отсутствует — пропущен")
            continue
        ws = wb[sheet]
        rows = []
        for row in ws.iter_rows(values_only=True):
            if all(v is None for v in row):
                continue
            rows.append([cell(v) for v in row])
        # обрезаем пустые хвостовые столбцы
        width = max((len(r) - next((i for i, v in enumerate(reversed(r)) if v != ""), len(r))) for r in rows)
        rows = [r[:width] for r in rows]
        with open(OUT / name, "w", newline="", encoding="utf-8-sig") as f:
            csv.writer(f, lineterminator="\n").writerows(rows)
        print(f"{sheet:12s} -> csv/{name}: {len(rows)} строк")

    # предупреждение: формулы без кэшированных значений
    wb_f = openpyxl.load_workbook(XLSX)
    for sheet in SHEETS:
        if sheet not in wb_f.sheetnames:
            continue
        for row_f, row_v in zip(wb_f[sheet].iter_rows(), wb[sheet].iter_rows()):
            for cf, cv in zip(row_f, row_v):
                if isinstance(cf.value, str) and cf.value.startswith("=") and cv.value is None:
                    missing_values += 1
    if missing_values:
        print(f"! {missing_values} формул без вычисленного значения — пересохраните файл и повторите экспорт")
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(main())
